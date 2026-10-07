"""
excel_formato.py
Crea el Excel "Control de Equipos" con el MISMO formato de ESTADO_TIENDAS_2026.xlsx
(encabezado institucional, grupos de equipos, lista de tiendas, listas desplegables,
colores condicionales BUENO / REGULAR / MALO / N/A y filas de totales).
"""
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as ci
from openpyxl.utils import get_column_letter as cl
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation

SHEET_NAME = "Control de Equipos"
FIRST_ROW = 8          # primera tienda
FONT_MAIN = "Avenir LT Pro 55 Roman"

# ----------------------------------------------------------------- colores
VERDE = "046A38"
VERDE_CLARO = "C2E2D2"
VERDE_SUB = "BFD7CB"
VERDE_INFO = "EAF2EE"
ZEBRA = "F4F8F6"
VERDE_TXT = "024021"

def _fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")

def _side(style, color):
    return Side(style=style, color=color)

THIN_SUB = _side("thin", VERDE_SUB)
THIN_GRIS = _side("thin", "D3D3D3")
MED = _side("medium", VERDE)
THIN_V = _side("thin", VERDE)

# ----------------------------------------------------------------- estructura
# Grupos de la fila 5 (título, col inicio, col fin)
GRUPOS = [
    ("POS I", "D", "O"), ("POS II", "P", "W"),
    ("PERIFÉRICOS POS 1", "X", "Z"), ("PERIFÉRICOS POS II", "AA", "AC"),
    ("IMPRESORA LÁSER", "AD", "AE"), ("PANTALLAS", "AF", "AI"),
    ("EQUIPO DE SONIDO", "AJ", "AK"), ("ADICIONALES", "AL", "AM"),
    ("EQUIPOS DE RED", "AN", "AQ"), ("MONITOREO", "AR", "AS"),
]

# Equipos de la fila 6 (título, col inicio, col fin). Si inicio == fin -> solo "Estado"
EQUIPOS = [
    ("Computador", "D", "E"), ("Teléfono", "F", "G"), ("Biométrico", "H", "I"),
    ("Lector de Código", "J", "K"), ("Cajón Monedero", "L", "M"), ("Impresora POS", "N", "O"),
    ("Computador", "P", "Q"), ("Lector de Código", "R", "S"),
    ("Cajón Monedero", "T", "U"), ("Impresora POS", "V", "W"),
    ("Teclado", "X", "X"), ("Mouse", "Y", "Y"), ("Pantalla", "Z", "Z"),
    ("Teclado", "AA", "AA"), ("Mouse", "AB", "AB"), ("Pantalla", "AC", "AC"),
    ("Impresora Láser", "AD", "AE"),
    ("Pantalla Publicitaria I ", "AF", "AG"), ("Pantalla Publicitaria II", "AH", "AI"),
    ("Equipo de Sonido", "AJ", "AK"),
    ("Cables (Poder/Vid/Red)", "AL", "AL"), ("PadMouse", "AM", "AM"),
    ("AP (Access Point)", "AN", "AO"), ("SWITCH", "AP", "AQ"),
    ("Monitor de Cámaras", "AR", "AS"),
]
OBS_INI, OBS_FIN = "AT", "AX"
LAST_COL = ci("AX")

# Columnas "Estado" (llevan lista desplegable, colores y conteo)
COLS_ESTADO = [e[2] for e in EQUIPOS]

# Primera columna de cada bloque -> borde izquierdo grueso
INICIO_BLOQUE = {"F", "H", "J", "L", "N", "P", "R", "T", "V", "X", "AA", "AD", "AF",
                 "AH", "AJ", "AL", "AN", "AR", "AT"}

ANCHOS = {
    "A": 13.43, "B": 48.43, "C": 20.86, "D": 15.43, "E": 13.43, "F": 29.29, "G": 13.71,
    "H": 28.43, "I": 15.71, "J": 17.0, "K": 11.57, "L": 17.29, "M": 12.0, "N": 14.0,
    "O": 11.57, "P": 17.86, "Q": 12.0, "R": 17.0, "S": 14.0, "T": 12.0, "U": 24.29,
    "V": 15.71, "W": 14.0, "X": 12.0, "Y": 14.0, "Z": 13.14, "AA": 12.0, "AB": 14.0,
    "AC": 13.14, "AD": 11.0, "AE": 12.0, "AF": 14.0, "AG": 12.0, "AH": 14.0, "AI": 12.0,
    "AJ": 16.0, "AK": 13.14, "AL": 24.0, "AM": 16.57, "AN": 23.71, "AO": 11.57,
    "AP": 15.0, "AQ": 11.57, "AR": 19.14, "AS": 11.57, "AT": 22.57, "AU": 16.57,
    "AV": 17.71, "AW": 18.71, "AX": 26.71,
}

# Centro de costo, tienda, ciudad
TIENDAS = [
    ("8B", "RIVIERA ARMENIA", "ARMENIA"),
    (13, "RIVIERA BUENAVISTA BARRANQUILLA", "BARRANQUILLA"),
    (7, "RIVIERA UNICENTRO 1-040", "BOGOTÁ"),
    (32, "RIVIERA PALATINO", "BOGOTÁ"),
    (34, "RIVIERA SANTA ANA", "BOGOTÁ"),
    (21, "INEDITO PERFUME DE AUTOR", "BOGOTÁ"),
    (36, "RIVIERA EL RETIRO", "BOGOTÁ"),
    (38, "RIVIERA SANTA FE", "BOGOTÁ"),
    (41, "RIVIERA GRAN ESTACION", "BOGOTÁ"),
    (8, "RIVIERA UNICENTRO 1-252", "BOGOTÁ"),
    (2, "RIVIERA SALITRE", "BOGOTÁ"),
    (83, "LA HORA TALLER", "BOGOTÁ"),
    (1, "RIVIERA ANDINO", "BOGOTÁ"),
    (16, "RIVIERA UNICENTRO CALI", "CALI"),
    (28, "RIVIERA CHIPICHAPE", "CALI"),
    (44, "RIVIERA CENTENARIO", "CALI"),
    (15, "RIVIERA CARTAGENA", "CARTAGENA"),
    ("2A", "RIVIERA CARIBE PLAZA", "CARTAGENA"),
    ("1E", "RIVIERA PLAZA BOCAGRANDE", "CARTAGENA"),
    (46, "RIVIERA VENTURA PLAZA", "CÚCUTA"),
    (43, "RIVIERA UNICENTRO CUCUTA", "CÚCUTA"),
    (35, "RIVIERA IBAGUE", "IBAGUÉ"),
    ("2F", "RIVIERA ACQUA IBAGUE", "IBAGUÉ"),
    (29, "RIVIERA MANIZALES", "MANIZALES"),
    (11, "RIVIERA TESORO MEDELLIN", "MEDELLÍN"),
    (12, "RIVIERA UNICENTRO MEDELLIN", "MEDELLÍN"),
    ("8A", "RIVIERA SANTA FE MEDELLIN", "MEDELLÍN"),
    ("2R", "RIVIERA MONTERIA ALAMEDA", "MONTERÍA"),
    ("8G", "RIVIERA SANTA LUCIA", "NEIVA"),
    ("3A", "RIVIERA UNICENTRO PEREIRA", "PEREIRA"),
    (47, "RIVIERA POPAYAN", "POPAYÁN"),
    (39, "RIVIERA SANTA MARTA", "SANTA MARTA"),
    (48, "RIVIERA VALLEDUPAR", "VALLEDUPAR"),
    (40, "RIVIERA VILLAVICENCIO", "VILLAVICENCIO"),
    ("7D", "RIVIERA YOPAL", "YOPAL"),
]
LAST_ROW = FIRST_ROW + len(TIENDAS) - 1          # 42
ROW_BUENO, ROW_REGULAR, ROW_MALO = LAST_ROW + 1, LAST_ROW + 2, LAST_ROW + 3

# Mapeo del formulario -> columnas del Excel, por número de caja.
# (columna_activo, columna_estado); None = esa columna no existe en el formato.
# Caja 1 -> POS I, Caja 2 -> POS II. Lo que no tenga columna va a OBSERVACIONES.
MAPA_CAJAS = {
    1: {"PC": ("D", "E"), "Telefono": ("F", "G"), "Biometrico": ("H", "I"),
        "Lector": ("J", "K"), "Cajon": ("L", "M"), "Impresora": ("N", "O"),
        "Teclado": (None, "X"), "Mouse": (None, "Y"), "Pantalla": (None, "Z")},
    2: {"PC": ("P", "Q"), "Lector": ("R", "S"), "Cajon": ("T", "U"),
        "Impresora": ("V", "W"),
        "Teclado": (None, "AA"), "Mouse": (None, "AB"), "Pantalla": (None, "AC")},
}
NOMBRES = {"PC": "Computador", "Telefono": "Teléfono", "Biometrico": "Biométrico",
           "Lector": "Lector", "Cajon": "Cajón", "Impresora": "Impresora POS",
           "Teclado": "Teclado", "Mouse": "Mouse", "Pantalla": "Pantalla"}


# ----------------------------------------------------------------- helpers
def _rango(ws, ini, fin, row_ini, row_fin=None):
    row_fin = row_fin or row_ini
    for r in range(row_ini, row_fin + 1):
        for c in range(ci(ini), ci(fin) + 1):
            yield ws.cell(r, c)


def _bordes(cell, left=None, right=None, top=None, bottom=None):
    b = cell.border
    cell.border = Border(left=left or b.left, right=right or b.right,
                         top=top or b.top, bottom=bottom or b.bottom)


def _encabezado_superior(ws):
    f_blanca = Font(name=FONT_MAIN, size=12, bold=True, color="FFFFFF")
    # Título (A1:B3 combinado + relleno hasta E3)
    for c in _rango(ws, "A", "E", 1, 3):
        c.fill = _fill(VERDE)
        c.font = f_blanca
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("A1:B3")
    ws["A1"] = ("SISTEMA DE GESTIÓN ADMINISTRATIVA\nFORMATO DE CONTROL Y SEGUIMIENTO\n"
                "ESTADO DE EQUIPOS TECNOLÓGICOS")

    hoy = datetime.now()
    etiquetas = {"F1": "FECHA DE EMISON", "F2": "VERSIÓN:", "F3": "FECHA REVISIÓN:",
                 "H1": "ELABORADO POR:", "H2": "REVISADO POR:", "H3": "APROBADO POR:"}
    valores = {"G1": hoy, "G2": "1.0.0", "G3": hoy.strftime("%d/%m/%Y"),
               "I1": "Cristian Arciniegas"}
    for r in (1, 2, 3):
        ws.merge_cells(f"I{r}:K{r}")
    for c in _rango(ws, "F", "K", 1, 3):
        c.fill = _fill(VERDE_INFO)
        color = "D3D3D3" if c.column <= 8 else VERDE_SUB
        s = _side("thin", color)
        c.border = Border(left=s, right=s, top=s, bottom=s)
        c.font = Font(name=FONT_MAIN, size=12, color="333333")
        c.alignment = Alignment(vertical="center")
    for coord, txt in etiquetas.items():
        ws[coord] = txt
        ws[coord].font = Font(name=FONT_MAIN, size=12, bold=True, color=VERDE)
        ws[coord].alignment = Alignment(horizontal="right", vertical="center")
    for coord, val in valores.items():
        ws[coord] = val
    ws["G1"].number_format = "dd/mm/yyyy"
    for coord in ("G1", "G2", "G3"):
        ws[coord].alignment = Alignment(horizontal="left", vertical="center")
    ws["I1"].alignment = Alignment(vertical="center", wrap_text=True)


def _encabezado_tabla(ws):
    f_grupo = Font(name=FONT_MAIN, size=12, bold=True, color="FFFFFF")
    f_equipo = Font(name=FONT_MAIN, size=12, bold=True, color=VERDE_TXT)
    f_sub = Font(name=FONT_MAIN, size=12, bold=True, color="444444")
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # --- A:C (C. COSTOS / TIENDA / CIUDAD), combinadas filas 5-7
    for col, txt in (("A", "C. COSTOS"), ("B", "TIENDA / SUCURSAL"), ("C", "CIUDAD")):
        ws.merge_cells(f"{col}5:{col}7")
        ws[f"{col}5"] = txt
        for r in (5, 6, 7):
            c = ws[f"{col}{r}"]
            c.fill = _fill(VERDE)
            c.font = f_grupo
            c.alignment = centro
            _bordes(c, left=MED if col in "AB" or col == "C" else None,
                    right=MED if col == "B" else THIN_V,
                    top=MED if r == 5 else None, bottom=MED if r == 7 else None)

    # --- Fila 5: grupos
    for titulo, ini, fin in GRUPOS:
        if ini != fin:
            ws.merge_cells(f"{ini}5:{fin}5")
        ws[f"{ini}5"] = titulo
        for c in _rango(ws, ini, fin, 5):
            c.fill = _fill(VERDE)
            c.font = f_grupo
            c.alignment = centro
            _bordes(c, top=MED, bottom=THIN_V)
        _bordes(ws[f"{ini}5"], left=MED)
        _bordes(ws[f"{fin}5"], right=MED)

    # --- Fila 6 y 7: equipos y Activo / Estado
    for titulo, ini, fin in EQUIPOS:
        if ini != fin:
            ws.merge_cells(f"{ini}6:{fin}6")
        ws[f"{ini}6"] = titulo
        for c in _rango(ws, ini, fin, 6):
            c.fill = _fill(VERDE_CLARO)
            c.font = f_equipo
            c.alignment = centro
            s = THIN_V
            c.border = Border(left=s, right=s, top=s, bottom=s)
        subs = ["Activo", "Estado"] if ini != fin else ["Estado"]
        for c, txt in zip(_rango(ws, ini, fin, 7), subs):
            c.value = txt
            c.fill = _fill(VERDE_SUB)
            c.font = f_sub
            c.alignment = Alignment(horizontal="center", vertical="center")
            s = THIN_V
            c.border = Border(left=s, right=s, top=s, bottom=s)
        if ini in INICIO_BLOQUE or ini == "D":
            _bordes(ws[f"{ini}5"], left=MED)
            for r in (6, 7):
                _bordes(ws[f"{ini}{r}"], left=MED if ini != "D" else None)

    # --- OBSERVACIONES (AT5:AX7)
    for c in _rango(ws, OBS_INI, OBS_FIN, 5, 7):
        c.fill = _fill(VERDE)
    ws.merge_cells(f"{OBS_INI}6:{OBS_FIN}7")
    ws[f"{OBS_INI}6"] = "OBSERVACIONES"
    ws[f"{OBS_INI}6"].font = f_grupo
    ws[f"{OBS_INI}6"].alignment = Alignment(horizontal="center", vertical="center")
    for c in _rango(ws, OBS_INI, OBS_FIN, 5, 7):
        _bordes(c, top=MED if c.row == 5 else None,
                left=MED if c.column == ci(OBS_INI) else None,
                right=MED if c.column == ci(OBS_FIN) else None)


def _cuerpo(ws):
    f_dato = Font(name=FONT_MAIN, size=12, color="333333")
    centro = Alignment(horizontal="center", vertical="center")
    for i, (cc, tienda, ciudad) in enumerate(TIENDAS):
        r = FIRST_ROW + i
        zebra = (r % 2 == 1)
        for col in range(1, ci("AS") + 1):
            c = ws.cell(r, col)
            c.font = f_dato
            c.alignment = centro
            if zebra:
                c.fill = _fill(ZEBRA)
            c.border = Border(left=THIN_SUB, right=THIN_SUB, top=THIN_SUB, bottom=THIN_SUB)
            if cl(col) in ("A", "B", "C") or cl(col) in INICIO_BLOQUE:
                _bordes(c, left=MED)
        _bordes(ws.cell(r, 2), right=MED)
        ws.cell(r, 1, cc)
        ws.cell(r, 2, tienda)
        ws.cell(r, 3, ciudad)

        # Observaciones (AT:AX combinadas por fila)
        ws.merge_cells(f"{OBS_INI}{r}:{OBS_FIN}{r}")
        for c in _rango(ws, OBS_INI, OBS_FIN, r):
            c.font = Font(name=FONT_MAIN, size=12, color="333333")
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if zebra:
                c.fill = _fill(ZEBRA)
            _bordes(c, top=THIN_SUB, bottom=THIN_SUB,
                    left=MED if c.column == ci(OBS_INI) else None,
                    right=MED if c.column == ci(OBS_FIN) else None)
    # cierre inferior de la tabla
    for c in _rango(ws, OBS_INI, OBS_FIN, LAST_ROW):
        _bordes(c, bottom=MED)


def _resumen(ws):
    estilos = [
        (ROW_BUENO, "Cantidad Equipos BUENO", "Bueno", "E6F4EA", "006100"),
        (ROW_REGULAR, "Cantidad Equipos REGULAR", "Regular", "FFF9E6", "9C6500"),
        (ROW_MALO, "Cantidad Equipos MALO", "Malo", "FCE8E6", "9C0006"),
    ]
    for r, etiqueta, palabra, bg, fg in estilos:
        ws.cell(r, 1).fill = _fill(ZEBRA)
        for col in range(1, ci("AS") + 1):
            c = ws.cell(r, col)
            _bordes(c, left=THIN_SUB, right=THIN_SUB, top=THIN_SUB, bottom=THIN_SUB)
            if col <= 2 or cl(col) in INICIO_BLOQUE:
                _bordes(c, left=MED)
            if r == ROW_MALO:
                _bordes(c, bottom=MED)
            if col >= 2:
                c.fill = _fill(bg)
                c.font = Font(name=FONT_MAIN, size=12, bold=True, color=fg)
                c.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 2, etiqueta)
        for col in range(4, ci("AS") + 1):
            letra = cl(col)
            if letra in COLS_ESTADO:
                ws.cell(r, col,
                        f'=COUNTIF({letra}{FIRST_ROW}:{letra}{LAST_ROW}, "{palabra}")')
            else:
                ws.cell(r, col, "-")


def _validaciones_y_colores(ws):
    dv = DataValidation(type="list", formula1='"BUENO,REGULAR,MALO,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    reglas = [
        ("BUENO", "C6EFCE", "006100"),
        ("REGULAR", "FFEB9C", "9C6500"),
        ("MALO", "FFC7CE", "9C0006"),
        ("N/A", "D9D9D9", None),
    ]
    for col in COLS_ESTADO:
        rango = f"{col}{FIRST_ROW}:{col}{LAST_ROW}"
        dv.add(rango)
        for texto, bg, fg in reglas:
            dxf = DifferentialStyle(
                font=Font(color=fg) if fg else None,
                fill=PatternFill(start_color=bg, end_color=bg, bgColor=bg, fill_type="solid"),
            )
            ws.conditional_formatting.add(rango, Rule(
                type="containsText", operator="containsText", text=texto, dxf=dxf,
                formula=[f'NOT(ISERROR(SEARCH("{texto}",{col}{FIRST_ROW})))'],
            ))


def crear_excel_formato(ruta):
    """Crea el libro completo con el formato institucional y lo guarda en `ruta`."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    for col, ancho in ANCHOS.items():
        ws.column_dimensions[col].width = ancho
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 20.1
    ws.row_dimensions[4].height = 9.95
    ws.row_dimensions[5].height = 15.75
    ws.row_dimensions[6].height = 33
    ws.row_dimensions[7].height = 18
    for r in range(FIRST_ROW, ROW_MALO + 1):
        ws.row_dimensions[r].height = 20.1

    _encabezado_superior(ws)
    _encabezado_tabla(ws)
    _cuerpo(ws)
    _resumen(ws)
    _validaciones_y_colores(ws)

    ws.freeze_panes = "D8"
    ws.sheet_view.zoomScale = 55
    ws.sheet_view.zoomScaleNormal = 55
    wb.save(ruta)
    return ruta


def guardar_registros(ruta, registros):
    """
    Escribe los registros del formulario en la fila de la tienda (columna B).
    Caja 1 -> POS I, Caja 2 -> POS II. Cajas 3-5 y los equipos sin columna
    (ej. teléfono de la caja 2) se anotan en OBSERVACIONES.
    Devuelve la lista de tiendas no encontradas.
    """
    wb = openpyxl.load_workbook(ruta)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    filas = {str(ws.cell(r, 2).value).strip().upper(): r
             for r in range(FIRST_ROW, LAST_ROW + 1)}
    no_encontradas = []

    for reg in registros:
        r = filas.get(str(reg.get("Oficina", "")).strip().upper())
        if r is None:
            no_encontradas.append(reg.get("Oficina"))
            continue

        num = int(str(reg.get("Numero_Caja", "Caja 1")).split()[-1])
        mapa = MAPA_CAJAS.get(num, {})
        notas = []

        for equipo, nombre in NOMBRES.items():
            activo = reg.get(f"{equipo}_Activo", reg.get(f"{equipo}_Activa", ""))
            estado = reg.get(f"{equipo}_Estado", "")
            if equipo in mapa:
                col_act, col_est = mapa[equipo]
                if col_act and activo:
                    ws[f"{col_act}{r}"] = activo
                if estado:
                    ws[f"{col_est}{r}"] = estado
            elif activo or estado:
                notas.append(f"{nombre}: {activo or '-'} ({estado or '-'})")

        if notas:
            texto = f"Caja {num} -> " + "; ".join(notas)
            celda = ws[f"{OBS_INI}{r}"]
            celda.value = f"{celda.value} | {texto}" if celda.value else texto

    wb.save(ruta)
    return no_encontradas
